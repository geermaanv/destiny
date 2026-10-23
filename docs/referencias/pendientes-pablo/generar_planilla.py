"""Planilla "Pendientes con Pablo": preguntas para responder y los textos de Tu día y Mi carta para revisar."""

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(__file__), "../../../apps/api/app/content/")
OUT = sys.argv[1]
carta = json.load(open(ROOT + "carta_v1.json"))
dia = json.load(open(ROOT + "transitos_v1.json"))

EDIT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="25225B")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

POINTS = {"sol": "Sol", "luna": "Luna", "asc": "Ascendente", "desc": "Descendente", "mc": "Medio Cielo", "ic": "Fondo del Cielo", "venus": "Venus", "saturno": "Saturno"}
DIMENSION = {"sol": "Identidad", "asc": "Identidad", "venus": "Amor", "desc": "Amor", "mc": "Vocación", "saturno": "Vocación", "luna": "Emociones", "ic": "Emociones"}
ASPECT = {
    "conjunción": "Conjunción (la Luna pasa por tu signo solar)",
    "sextil": "Sextil (se complementan)",
    "cuadratura": "Cuadratura (tensión)",
    "trígono": "Trígono (fluye)",
    "oposición": "Oposición (polos opuestos)",
    "sin aspecto mayor": "Sin aspecto mayor (día tranquilo)",
}
STYLE = {"directo": "Directo (Mercurio en fuego)", "practico": "Práctico (tierra)", "analitico": "Analítico (aire)", "intuitivo": "Intuitivo (agua)"}


def header(ws, values, widths):
    for i, (v, w) in enumerate(zip(values, widths), start=1):
        c = ws.cell(row=1, column=i, value=v)
        c.fill, c.font, c.alignment = HEAD, WHITE, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def review_columns(ws, first_row, last_row, ok_col):
    dv = DataValidation(type="list", formula1='"OK,Corregir"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(first_row, last_row + 1):
        for c in (ok_col, ok_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.alignment = EDIT, WRAP
        dv.add(ws.cell(row=r, column=ok_col))


wb = Workbook()

# --- Pendientes ---
ws = wb.active
ws.title = "Pendientes"
header(ws, ["#", "Tema", "Pregunta", "Opciones", "Si no respondés", "Tu respuesta", "Comentario"], [4, 18, 46, 46, 30, 16, 40])
QUESTIONS = [
    ("1", "Test", "Alguien que no hizo el test de personalidad", "A) El número global usa solo la astrología\nB) No aparece cuando se filtra por test", "A (como funciona hoy)", ["A", "B"]),
    ("2a", "Numerología", "¿Qué te parece la propuesta (planilla que te compartí)?", "A) Me cierra, avancemos\nB) La ajusto en la planilla\nC) Prefiero otro método, lo charlamos", "—", ["A", "B", "C"]),
    ("2b", "Numerología", "¿Usamos el nombre completo de nacimiento?", "A) Imprescindible\nB) Opcional, para quien quiera profundizar\nC) No hace falta, alcanza con la fecha", "Solo la fecha; el nombre opcional más adelante", ["A", "B", "C"]),
    ("3", "Test", "Título del grupo de atracción del test", 'A) "Química y atracción"\nB) Dejarlo como "Atracción y sexualidad"', "B (queda como está)", ["A", "B"]),
    ("4", "Indagar", "¿Las respuestas de atracción del test se usan para responder en Indagar?", "A) No, nunca (son privadas)\nB) Sí, sin mostrarlas", "A", ["A", "B"]),
    ("5", "Fotos", "Fotos por elemento: 4 fotos (Fuego, Tierra, Aire, Agua) y Destiny elige cuál mostrar según el día", "A) Me gusta, hagámoslo\nB) No, dejemos las 4 fotos comunes", "B", ["A", "B"]),
    ("6", "Cobrar", "¿Seguimos sin cobrar nada durante los primeros 6 meses?", "A) Confirmo\nB) Quiero charlarlo", "A", ["A", "B"]),
    ("7", "Dominio", "Nombre del dominio para salir a producción (y si te importa la terminación: .com, .app, .com.ar, .ai…)", "A) Ya tengo uno: escribilo en Comentario\nB) Pasame opciones disponibles", "B", ["A", "B"]),
    ("8", "Textos", "Revisar los textos de Tu día y de Mi carta (hojas siguientes): marcá OK o Corregir y escribí la corrección", "—", "—", None),
    ("9", "Estilo", "¿Qué te parece cómo te habla la app y los colores según tu carta?", "Contame en Comentario", "—", None),
    ("10", "Probar", "Probar lo nuevo (al entrar al test te pide aceptar un aviso de privacidad)", "—", "—", None),
]
for r, (n, theme, q, opts, default, choices) in enumerate(QUESTIONS, start=2):
    for c, v in enumerate([n, theme, q, opts, default], start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = WRAP
    ws.cell(row=r, column=2).font = BOLD
    for c in (6, 7):
        cell = ws.cell(row=r, column=c)
        cell.fill, cell.alignment = EDIT, WRAP
    if choices:
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(ws.cell(row=r, column=6))
note = len(QUESTIONS) + 3
ws.cell(row=note, column=2, value="Cómo usar esta planilla").font = BOLD
ws.cell(row=note + 1, column=2, value="Completá lo amarillo. Podés responder solo algunas: las que no contestes quedan como dice \"Si no respondés\".")
ws.cell(row=note + 2, column=2, value="Los textos de las hojas siguientes los redactó la IA a partir de la astrología; corregí lo que no te suene bien.")

# --- Tu día ---
ws = wb.create_sheet("Tu día")
header(ws, ["Aspecto de la Luna con tu Sol", "Signo de la Luna", "Título", "Texto", "¿OK?", "Corrección"], [34, 14, 26, 70, 10, 50])
r = 2
for aspect, signs in dia["textos"].items():
    for sign, entry in signs.items():
        for c, v in enumerate([ASPECT.get(aspect, aspect), sign, entry["titulo"], entry["texto"]], start=1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        r += 1
review_columns(ws, 2, r - 1, 5)

# --- Tu día por estilo ---
ws = wb.create_sheet("Tu día por estilo")
header(ws, ["Estilo de redacción", "Aspecto de la Luna con tu Sol", "Signo de la Luna", "Título", "Texto", "¿OK?", "Corrección"], [26, 34, 14, 26, 70, 10, 50])
r = 2
for style, aspects in dia["por_registro"].items():
    for aspect, signs in aspects.items():
        for sign, entry in signs.items():
            for c, v in enumerate([STYLE.get(style, style), ASPECT.get(aspect, aspect), sign, entry["titulo"], entry["texto"]], start=1):
                ws.cell(row=r, column=c, value=v).alignment = WRAP
            r += 1
review_columns(ws, 2, r - 1, 6)

# --- Mi carta: signos ---
ws = wb.create_sheet("Mi carta · signos")
header(ws, ["Área", "Punto", "Signo", "Texto", "¿OK?", "Corrección"], [14, 18, 14, 80, 10, 50])
r = 2
for point, signs in carta["signo"].items():
    for sign, text in signs.items():
        for c, v in enumerate([DIMENSION.get(point, ""), POINTS.get(point, point), sign, text], start=1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        r += 1
review_columns(ws, 2, r - 1, 5)

# --- Mi carta: casas ---
ws = wb.create_sheet("Mi carta · casas")
header(ws, ["Área", "Punto", "Casa", "Texto", "¿OK?", "Corrección"], [14, 18, 8, 80, 10, 50])
r = 2
for point, houses in carta["casa"].items():
    for house, text in houses.items():
        for c, v in enumerate([DIMENSION.get(point, ""), POINTS.get(point, point), int(house), text], start=1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
        r += 1
review_columns(ws, 2, r - 1, 5)

wb.save(OUT)
print("ok", {s.title: s.max_row - 1 for s in wb.worksheets})
