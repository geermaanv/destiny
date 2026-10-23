"""Genera la planilla de la propuesta de numerología v0 para Pablo (Destiny).

Uso (necesita openpyxl, que no es dependencia del proyecto):
    python3 -m venv /tmp/xl && /tmp/xl/bin/pip install openpyxl
    /tmp/xl/bin/python docs/referencias/numerologia/generar_planilla.py docs/referencias/numerologia/Numerologia-Destiny-propuesta-v0.xlsx


Todo editable en celdas amarillas; la hoja Calculadora usa fórmulas, así al cambiar una tabla se ve
el resultado al instante.
"""

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]

AXES = ["Química", "Afecto", "Comunicación", "Compromiso"]
NUMS = list(range(1, 10))
MEANING = {
    1: "Iniciativa, independencia, liderazgo",
    2: "Cooperación, sensibilidad, vínculo",
    3: "Expresión, creatividad, alegría",
    4: "Estructura, constancia, trabajo",
    5: "Libertad, cambio, aventura",
    6: "Cuidado, familia, responsabilidad",
    7: "Introspección, análisis, búsqueda",
    8: "Ambición, logro, poder",
    9: "Compasión, idealismo, entrega",
}
MASTERS = {11: (2, "Intuición, inspiración (se lee como 2)"), 22: (4, "Constructor (se lee como 4)"), 33: (6, "Maestro, servicio (se lee como 6)")}
# Afinidad de cada número con cada eje (0-100): qué aporta a ese aspecto del vínculo.
AFFINITY = {
    1: [70, 40, 60, 50],
    2: [50, 90, 60, 80],
    3: [70, 60, 90, 40],
    4: [40, 60, 50, 90],
    5: [90, 40, 80, 30],
    6: [50, 90, 60, 90],
    7: [40, 50, 70, 60],
    8: [80, 50, 50, 80],
    9: [60, 80, 70, 60],
}
# Compatibilidad base entre caminos de vida (simétrica). Grupos clásicos: 1-5-7 (mentales),
# 3-6-9 (creativos/emocionales), 2-4-8 (prácticos); combinaciones clásicas en tensión: 4-5, 1-8.
UPPER = {
    1: {1: 65, 2: 60, 3: 85, 4: 50, 5: 80, 6: 55, 7: 75, 8: 50, 9: 80},
    2: {2: 70, 3: 60, 4: 85, 5: 40, 6: 90, 7: 55, 8: 85, 9: 65},
    3: {3: 70, 4: 45, 5: 85, 6: 85, 7: 55, 8: 50, 9: 90},
    4: {4: 75, 5: 35, 6: 80, 7: 70, 8: 90, 9: 50},
    5: {5: 65, 6: 45, 7: 80, 8: 55, 9: 70},
    6: {6: 75, 7: 45, 8: 70, 9: 90},
    7: {7: 70, 8: 50, 9: 65},
    8: {8: 65, 9: 55},
    9: {9: 75},
}
BASE = {(a, b): UPPER[min(a, b)][max(a, b)] for a in NUMS for b in NUMS}
# Pesos por vínculo: los mismos del motor Destiny 1.0 de Pablo (synastry_v2.json).
WEIGHTS = {
    "Pareja": [0.25, 0.25, 0.25, 0.25],
    "Amistad": [0.10, 0.35, 0.40, 0.15],
    "Trabajo": [0.10, 0.10, 0.40, 0.40],
    "Casual": [0.70, 0.10, 0.15, 0.05],
}

EDIT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="25225B")
OUTF = PatternFill("solid", fgColor="E2EFDA")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Border(*(Side(style="thin", color="DCD9EC"),) * 4)


def header(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.fill, c.font, c.alignment = HEAD, WHITE, Alignment(horizontal="center", wrap_text=True)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# --- Leeme ---
ws = wb.active
ws.title = "Leeme"
lines = [
    ("Destiny · Numerología · Propuesta v0 (borrador para que Pablo revise y corrija)", BOLD),
    ("",),
    ("Qué es: un modelo para sumar la numerología como tercer sistema de compatibilidad, junto a la astrología y el test.",),
    ("Usa los mismos 4 ejes (Química, Afecto, Comunicación, Compromiso) y los mismos pesos por vínculo que tu motor 1.0,",),
    ("así entra en Círculo y en el número global sin cambiar las pantallas.",),
    ("",),
    ("Qué se calcula (solo con la fecha de nacimiento, sin pedir el nombre completo):", BOLD),
    ("• Camino de vida: día, mes y año reducidos cada uno a un dígito, sumados y reducidos (se conservan 11, 22 y 33).",),
    ("• Número del día: el día de nacimiento reducido (se conservan 11 y 22).",),
    ("• Para comparar, los maestros se leen como su base: 11 → 2, 22 → 4, 33 → 6.",),
    ("",),
    ("Cómo se compara a dos personas:", BOLD),
    ("1. Matriz base: compatibilidad de dos números (hoja Matriz).",),
    ("2. Cada eje ajusta la base según cuánto aportan los dos números a ese eje (hoja Números, afinidad por eje):",),
    ("   eje = base + Ajuste × (promedio de afinidades − 50) / 50, entre 0 y 100.",),
    ("3. Se combina Camino de vida (70 %) y Número del día (30 %) (hoja Parámetros).",),
    ("4. Total de cada vínculo = suma de ejes × pesos del vínculo (los de tu motor).",),
    ("",),
    ("Qué te pido:", BOLD),
    ("• Todas las celdas amarillas son editables: matriz, afinidades, pesos y parámetros.",),
    ("• Probá con la hoja Calculadora (poné dos fechas) y ajustá hasta que los resultados te cierren.",),
    ("• Si preferís otro método (otros números, otra forma de comparar), contame y lo armo.",),
    ("",),
    ("Opcional, con el nombre completo de nacimiento (hoja Con nombre):", BOLD),
    ("• Expresión (todas las letras), Alma (vocales) y Personalidad (consonantes), con la tabla pitagórica (A=1 … I=9, J=1 …).",),
    ("• Alma → Afecto, Personalidad → Química, Expresión → Comunicación y Compromiso.",),
    ("• Se mezcla con lo de la fecha según el Peso del nombre (hoja Parámetros, 50 % de arranque).",),
    ("• Sería opcional en la app: el nombre que guardamos hoy es el apodo del perfil, que no sirve para esto.",),
    ("",),
    ("Los valores de este borrador salen de las reglas clásicas de la numerología (grupos 1-5-7, 3-6-9, 2-4-8 y",),
    ("combinaciones en tensión como 4-5 y 1-8). Son un punto de partida: el modelo final lo definís vos.",),
]
for i, line in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=line[0])
    if len(line) > 1:
        c.font = line[1]
ws.column_dimensions["A"].width = 120

# --- Números ---
ws = wb.create_sheet("Números")
header(ws, 1, ["Número", "Significado", *[f"Afinidad {a}" for a in AXES]])
for r, n in enumerate(NUMS, start=2):
    ws.cell(row=r, column=1, value=n).font = BOLD
    ws.cell(row=r, column=2, value=MEANING[n])
    for j, v in enumerate(AFFINITY[n]):
        c = ws.cell(row=r, column=3 + j, value=v)
        c.fill, c.border = EDIT, THIN
r = 12
ws.cell(row=r, column=1, value="Maestros").font = BOLD
for i, (m, (base, text)) in enumerate(MASTERS.items(), start=1):
    ws.cell(row=r + i, column=1, value=m)
    ws.cell(row=r + i, column=2, value=text)
widths(ws, [10, 42, 16, 16, 16, 16])

# --- Matriz ---
ws = wb.create_sheet("Matriz")
ws.cell(row=1, column=1, value="Compatibilidad base (0-100). Mantenerla simétrica: (a,b) = (b,a).").font = BOLD
header(ws, 2, ["", *NUMS])
for i, a in enumerate(NUMS):
    ws.cell(row=3 + i, column=1, value=a).font = BOLD
    for j, b in enumerate(NUMS):
        c = ws.cell(row=3 + i, column=2 + j, value=BASE[(a, b)])
        c.fill, c.border, c.alignment = EDIT, THIN, Alignment(horizontal="center")
widths(ws, [8] + [7] * 9)

# --- Parámetros ---
ws = wb.create_sheet("Parámetros")
params = [
    ("Peso Camino de vida", 0.7),
    ("Peso Número del día", 0.3),
    ("Ajuste por eje (puntos máximos)", 10),
    ("Peso del nombre, si se carga (hoja Con nombre)", 0.5),
]
for i, (k, v) in enumerate(params, start=1):
    ws.cell(row=i, column=1, value=k)
    c = ws.cell(row=i, column=2, value=v)
    c.fill, c.border = EDIT, THIN
header(ws, 5, ["Vínculo", *AXES])
for i, (k, ws_) in enumerate(WEIGHTS.items(), start=6):
    ws.cell(row=i, column=1, value=k).font = BOLD
    for j, v in enumerate(ws_):
        c = ws.cell(row=i, column=2 + j, value=v)
        c.fill, c.border, c.number_format = EDIT, THIN, "0%"
ws.cell(row=11, column=1, value="Los pesos por vínculo son los de tu motor Destiny 1.0 (cada fila suma 100 %).")
widths(ws, [34, 14, 14, 16, 14])


def reduce_formulas(ws, row, col, src, keep):
    """Suma de dígitos repetida, en celdas auxiliares, con aritmética simple (funciona en Excel,
    Numbers y Google Sheets). Devuelve la celda con el número reducido. keep: maestros a conservar."""
    d = lambda r: f"(INT({r}/1000)+MOD(INT({r}/100),10)+MOD(INT({r}/10),10)+MOD({r},10))"  # noqa: E731
    a = ws.cell(row=row, column=col, value=f"={d(src)}")
    b = ws.cell(row=row, column=col + 1, value=f"={d(a.coordinate)}")
    keep_src = "OR(" + ",".join(f"{src}={m}" for m in keep) + ")" if keep else "FALSE"
    keep_a = "OR(" + ",".join(f"{a.coordinate}={m}" for m in keep) + ")" if keep else "FALSE"
    keep_b = "OR(" + ",".join(f"{b.coordinate}={m}" for m in keep) + ")" if keep else "FALSE"
    out = ws.cell(
        row=row,
        column=col + 2,
        value=(
            f"=IF({src}<10,{src},IF({keep_src},{src},IF({a.coordinate}<10,{a.coordinate},IF({keep_a},{a.coordinate},"
            f"IF({b.coordinate}<10,{b.coordinate},IF({keep_b},{b.coordinate},{d(b.coordinate)}))))))"
        ),
    )
    return out


# --- Calculadora ---
ws = wb.create_sheet("Calculadora")
ws.cell(row=1, column=1, value="Poné dos fechas de nacimiento (celdas amarillas) y mirá la compatibilidad numerológica.").font = BOLD
header(ws, 3, ["", "Persona A", "Persona B"])
ws.cell(row=4, column=1, value="Fecha de nacimiento")
for col, value in [(2, "1990-05-20"), (3, "1988-11-07")]:
    from datetime import date

    c = ws.cell(row=4, column=col, value=date.fromisoformat(value))
    c.fill, c.border, c.number_format = EDIT, THIN, "dd/mm/yyyy"
labels = ["Día", "Mes", "Año", "Día reducido", "Mes reducido", "Año reducido", "Camino de vida", "Número del día", "Camino (base)", "Día (base)"]
for i, l in enumerate(labels, start=5):
    ws.cell(row=i, column=1, value=l)
aux_col = {2: 10, 3: 16}  # celdas auxiliares fuera de la vista
for col in (2, 3):
    L = get_column_letter(col)
    ws.cell(row=5, column=col, value=f"=DAY({L}4)")
    ws.cell(row=6, column=col, value=f"=MONTH({L}4)")
    ws.cell(row=7, column=col, value=f"=YEAR({L}4)")
    ax = aux_col[col]
    rd = reduce_formulas(ws, 5, ax, f"{L}5", [11, 22])
    rm = reduce_formulas(ws, 6, ax, f"{L}6", [11])
    ry = reduce_formulas(ws, 7, ax, f"{L}7", [])
    ws.cell(row=8, column=col, value=f"={rd.coordinate}")
    ws.cell(row=9, column=col, value=f"={rm.coordinate}")
    ws.cell(row=10, column=col, value=f"={ry.coordinate}")
    ws.cell(row=20, column=ax, value=f"={L}8+{L}9+{L}10")
    life = reduce_formulas(ws, 21, ax, get_column_letter(ax) + "20", [11, 22, 33])
    ws.cell(row=11, column=col, value=f"={life.coordinate}").font = BOLD
    day = reduce_formulas(ws, 22, ax, f"{L}5", [11, 22])
    ws.cell(row=12, column=col, value=f"={day.coordinate}").font = BOLD
    base = lambda ref: f"IF({ref}=11,2,IF({ref}=22,4,IF({ref}=33,6,{ref})))"  # noqa: E731
    ws.cell(row=13, column=col, value=f"={base(f'{L}11')}")
    ws.cell(row=14, column=col, value=f"={base(f'{L}12')}")
for c in range(10, 22):
    ws.column_dimensions[get_column_letter(c)].hidden = True

header(ws, 16, ["Eje", "Camino de vida", "Número del día", "Eje combinado"])


def axis_formula(a_ref, b_ref, axis_idx):
    m = f"INDEX(Matriz!$B$3:$J$11,{a_ref},{b_ref})"
    aff_col = get_column_letter(3 + axis_idx)
    affs = f"(INDEX(Números!${aff_col}$2:${aff_col}$10,{a_ref})+INDEX(Números!${aff_col}$2:${aff_col}$10,{b_ref}))/2"
    return f"MAX(0,MIN(100,{m}+Parámetros!$B$3*({affs}-50)/50))"


for i, axis in enumerate(AXES):
    r = 17 + i
    ws.cell(row=r, column=1, value=axis).font = BOLD
    ws.cell(row=r, column=2, value="=" + axis_formula("$B$13", "$C$13", i)).number_format = "0"
    ws.cell(row=r, column=3, value="=" + axis_formula("$B$14", "$C$14", i)).number_format = "0"
    c = ws.cell(row=r, column=4, value=f"=Parámetros!$B$1*B{r}+Parámetros!$B$2*C{r}")
    c.number_format, c.fill, c.font = "0", OUTF, BOLD

header(ws, 23, ["Vínculo", "Compatibilidad"])
for i, k in enumerate(WEIGHTS, start=24):
    ws.cell(row=i, column=1, value=k).font = BOLD
    wr = 6 + list(WEIGHTS).index(k)
    c = ws.cell(row=i, column=2, value=f"=SUMPRODUCT($D$17:$D$20,TRANSPOSE(Parámetros!$B${wr}:$E${wr}))")
    c.value = f"=D17*Parámetros!B{wr}+D18*Parámetros!C{wr}+D19*Parámetros!D{wr}+D20*Parámetros!E{wr}"
    c.number_format, c.fill, c.font = "0", OUTF, BOLD
ws.cell(row=29, column=1, value="Global numerológico (promedio de los 4 vínculos)").font = BOLD
c = ws.cell(row=29, column=2, value="=AVERAGE(B24:B27)")
c.number_format, c.fill, c.font = "0", OUTF, BOLD
widths(ws, [44, 18, 18, 18])

# --- Con nombre (opcional) ---
ws = wb.create_sheet("Con nombre")
ws.cell(row=1, column=1, value="Opcional: con el nombre completo de nacimiento (el del acta). Poné los nombres en las celdas amarillas.").font = BOLD
ws.cell(row=2, column=1, value="Las fechas son las de la hoja Calculadora. Tildes y ñ se leen como la letra sin tilde y N; la Y cuenta como consonante.")
header(ws, 3, ["", "Persona A", "Persona B"])
ws.cell(row=4, column=1, value="Nombre completo de nacimiento")
for col, name in [(2, "María Laura Gómez"), (3, "Juan Pablo Ruiz")]:
    c = ws.cell(row=4, column=col, value=name)
    c.fill, c.border = EDIT, THIN
NAME_AUX = {2: 10, 3: 16}


def letters(ref, vowels_only=False):
    arr = f'MID({ref},ROW(INDIRECT("1:"&LEN({ref}))),1)'
    vowel = f'*ISNUMBER(SEARCH({arr},"AEIOU"))' if vowels_only else ""
    return f"IF(LEN({ref})=0,0,SUMPRODUCT((CODE({arr})>=65)*(CODE({arr})<=90)*(MOD(CODE({arr})-65,9)+1){vowel}))"


rows = ["Expresión", "Alma", "Personalidad", "Expresión (base)", "Alma (base)", "Personalidad (base)"]
for i, label in enumerate(rows, start=5):
    ws.cell(row=i, column=1, value=label)
for col in (2, 3):
    L = get_column_letter(col)
    ax = NAME_AUX[col]
    norm = ws.cell(row=30, column=ax)
    up = f"UPPER({L}4)"
    for src, dst in [("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("Ü", "U"), ("Ñ", "N")]:
        up = f'SUBSTITUTE({up},"{src}","{dst}")'
    norm.value = "=" + up
    n = norm.coordinate
    raw_e = ws.cell(row=31, column=ax, value="=" + letters(n))
    raw_a = ws.cell(row=32, column=ax, value="=" + letters(n, vowels_only=True))
    raw_p = ws.cell(row=33, column=ax, value=f"={raw_e.coordinate}-{raw_a.coordinate}")
    for r, raw in [(5, raw_e), (6, raw_a), (7, raw_p)]:
        red = reduce_formulas(ws, 30 + r, ax, raw.coordinate, [11, 22, 33])
        ws.cell(row=r, column=col, value=f"={red.coordinate}").font = BOLD
        ws.cell(row=r + 3, column=col, value=f"=IF({L}{r}=11,2,IF({L}{r}=22,4,IF({L}{r}=33,6,{L}{r})))")
for c in range(10, 22):
    ws.column_dimensions[get_column_letter(c)].hidden = True

header(ws, 12, ["Eje", "Con la fecha", "Con el nombre", "Combinado"])
NAME_SRC = {"Química": 10, "Afecto": 9, "Comunicación": 8, "Compromiso": 8}  # fila del número base que usa cada eje
for i, axis in enumerate(AXES):
    r = 13 + i
    ws.cell(row=r, column=1, value=f"{axis} ({ {10: 'Personalidad', 9: 'Alma', 8: 'Expresión'}[NAME_SRC[axis]] })").font = BOLD
    ws.cell(row=r, column=2, value=f"=Calculadora!D{17 + i}").number_format = "0"
    ws.cell(row=r, column=3, value="=" + axis_formula(f"$B${NAME_SRC[axis]}", f"$C${NAME_SRC[axis]}", i)).number_format = "0"
    c = ws.cell(row=r, column=4, value=f"=(1-Parámetros!$B$4)*B{r}+Parámetros!$B$4*C{r}")
    c.number_format, c.fill, c.font = "0", OUTF, BOLD
header(ws, 18, ["Vínculo", "Solo fecha", "Con nombre"])
for i, k in enumerate(WEIGHTS, start=19):
    wr = 6 + list(WEIGHTS).index(k)
    ws.cell(row=i, column=1, value=k).font = BOLD
    ws.cell(row=i, column=2, value=f"=Calculadora!B{24 + list(WEIGHTS).index(k)}").number_format = "0"
    c = ws.cell(row=i, column=3, value=f"=D13*Parámetros!B{wr}+D14*Parámetros!C{wr}+D15*Parámetros!D{wr}+D16*Parámetros!E{wr}")
    c.number_format, c.fill, c.font = "0", OUTF, BOLD
ws.cell(row=24, column=1, value="Global (promedio de los 4 vínculos)").font = BOLD
ws.cell(row=24, column=2, value="=Calculadora!B29").number_format = "0"
c = ws.cell(row=24, column=3, value="=AVERAGE(C19:C22)")
c.number_format, c.fill, c.font = "0", OUTF, BOLD
widths(ws, [40, 22, 22, 16])

wb.save(OUT)
print("ok", OUT)
